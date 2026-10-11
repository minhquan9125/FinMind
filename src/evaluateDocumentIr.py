"""Compare financial IR cells with a keyed gold set or Poppler PDF words.

Usage: python src/evaluateDocumentIr.py report.pdf --ir-dir path/to/annual_ir \
    --gold-xlsx Gold_Set/gold_fpt_2024.xlsx --gold-sheet BCTN_2024 --gold-only
Poppler measures tokenization agreement only. Spreadsheet agreement becomes
gold_accuracy only after its rows have been checked against PDF images; the
result records whether that check was self-performed or independent.
"""

from __future__ import annotations

import argparse
import csv
from datetime import date, datetime
import json
import re
import subprocess
import sys
import unicodedata
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Any

from financialContent import expand_page_ranges
from util import hash_file, required_command, save_json


def _poppler_words(pdf: Path, number: int) -> list[dict[str, Any]]:
    result = subprocess.run([required_command("pdftotext"), "-f", str(number), "-l", str(number),
                             "-bbox-layout", str(pdf), "-"], capture_output=True, check=True)
    root = ET.fromstring(result.stdout)
    words: list[dict[str, Any]] = []
    for element in root.iter():
        if element.tag.rsplit("}", 1)[-1] != "word":
            continue
        words.append({"text": "".join(element.itertext()),
                      "bbox": [float(element.attrib[name]) for name in
                               ("xMin", "yMin", "xMax", "yMax")]})
    return words


def _value(text: str) -> str:
    return re.sub(r"\s+", "", unicodedata.normalize("NFC", text)).replace("−", "-")


def _observed_cells(table_dir: Path, index: dict[str, Any]) -> dict[tuple[str, str, str], list[dict[str, Any]]]:
    """Index IR cells by statement, printed row code and column year."""
    observed: dict[tuple[str, str, str], list[dict[str, Any]]] = {}
    for number in expand_page_ranges(index["structured_table_ranges"]):
        page = json.loads((table_dir / f"p-{number:03d}.ir.json").read_text(encoding="utf-8"))
        page_text = " ".join(block.get("text", "") for block in page["blocks"]
                             if block["type"] in ("heading", "paragraph", "form_code"))
        day_month = re.search(r"(?:ngày|ngay)\s+(\d{1,2})\s+tháng\s+(\d{1,2})\b",
                              page_text, flags=re.IGNORECASE)
        for block in page["blocks"]:
            if block["type"] != "table":
                continue
            statement = block.get("logical_table_id") or page["page_class"]
            periods = {column["key"]: re.findall(r"\b20\d{2}\b", " ".join(column["header_lines"]))
                       for column in block["columns"]}
            for row in block["rows"]:
                if not row.get("code"):
                    continue
                for column, cell in row["cells"].items():
                    if not periods.get(column):
                        continue
                    year = periods[column][-1]
                    period_end = None
                    if day_month:
                        try:
                            period_end = date(int(year), int(day_month[2]), int(day_month[1])).isoformat()
                        except ValueError:
                            pass
                    key = statement, str(row["code"]), year
                    observed.setdefault(key, []).append({
                        "page": number, "raw": cell.get("raw"),
                        "period_end": period_end,
                        "period_type": "instant" if statement == "balance_sheet" else "duration",
                    })
    return observed


def _gold_rows(gold_path: Path, sheet_name: str | None) -> list[dict[str, str]]:
    """Read a CSV or one explicitly named Excel sheet without changing the source."""
    if gold_path.suffix.lower() == ".csv":
        with gold_path.open("r", encoding="utf-8-sig", newline="") as stream:
            reader = csv.DictReader(stream)
            rows = list(reader)
            headers = reader.fieldnames or []
    elif gold_path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        from openpyxl.utils.datetime import from_excel

        workbook = load_workbook(gold_path, read_only=True, data_only=True)
        try:
            target_sheet = sheet_name or workbook.sheetnames[0]
            if target_sheet not in workbook.sheetnames:
                raise ValueError(f"Excel không có sheet {target_sheet!r}; có: {workbook.sheetnames}")
            values = workbook[target_sheet].values
            headers = [str(value).strip() if value is not None else "" for value in next(values)]
            rows = []
            for cells in values:
                if not any(value is not None for value in cells):
                    continue
                item: dict[str, str] = {}
                for header, value in zip(headers, cells):
                    if header == "period_end" and isinstance(value, (int, float)):
                        value = from_excel(value, epoch=workbook.epoch)
                    if isinstance(value, (date, datetime)):
                        value = value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
                    item[header] = "" if value is None else str(value).strip()
                rows.append(item)
        finally:
            workbook.close()
    else:
        raise ValueError("Gold set phải là file .csv hoặc .xlsx")
    required = {"statement", "code", "period_end", "period_type", "value_raw", "pdf_page"}
    if not required.issubset(headers):
        raise ValueError(f"Gold set thiếu cột: {sorted(required - set(headers))}")
    if not rows:
        raise ValueError("Gold set không có dòng dữ liệu")
    return rows


def _spreadsheet_gold_result(table_dir: Path, index: dict[str, Any], gold_path: Path,
                             sheet_name: str | None = None) -> dict[str, Any]:
    """Measure row/column agreement without claiming unreviewed labels as gold."""
    gold_rows = _gold_rows(gold_path, sheet_name)
    printed_pages = {item["pdf_page"]: item.get("printed_page")
                     for item in index["page_index"]}

    observed = _observed_cells(table_dir, index)
    seen: set[tuple[str, str, str, str]] = set()
    reference_pages: dict[tuple[str, str, str, str], int] = {}
    matched_ir_pages: dict[tuple[str, str, str, str], int] = {}
    details: list[dict[str, Any]] = []
    for csv_line, item in enumerate(gold_rows, start=2):
        try:
            period = date.fromisoformat((item["period_end"] or "").strip())
            page_number = int((item["pdf_page"] or "").strip())
        except ValueError as exc:
            raise ValueError(f"CSV dòng {csv_line}: period_end hoặc pdf_page không hợp lệ") from exc
        statement = (item["statement"] or "").strip()
        code = (item["code"] or "").strip()
        period_type = (item["period_type"] or "").strip()
        if statement not in ("balance_sheet", "income_statement", "cash_flow"):
            raise ValueError(f"CSV dòng {csv_line}: statement không hợp lệ: {statement!r}")
        expected_type = "instant" if statement == "balance_sheet" else "duration"
        if period_type != expected_type or not code or page_number < 1:
            raise ValueError(f"CSV dòng {csv_line}: code, period_type hoặc pdf_page không hợp lệ")
        key = statement, code, period.isoformat(), period_type
        if key in seen:
            raise ValueError(f"CSV gold trùng khóa {key} tại dòng {csv_line}")
        seen.add(key)
        reference_pages[key] = page_number

        raw = (item["value_raw"] or "").strip()
        candidates = observed.get((statement, code, str(period.year)), [])
        page_candidates = [candidate for candidate in candidates
                           if candidate["page"] == page_number
                           or str(printed_pages.get(candidate["page"])) == str(page_number)]
        exact_candidates = [candidate for candidate in page_candidates
                            if candidate["period_end"] == period.isoformat()
                            and candidate["period_type"] == period_type]
        if raw.casefold() in ("", "nan", "null", "none", "#ref!", "=#ref!"):
            status = "invalid_reference"
        elif not candidates:
            status = "missing"
        elif not page_candidates:
            status = "wrong_page"
        elif any(candidate["period_end"] is None for candidate in page_candidates):
            status = "period_unresolved"
        elif not exact_candidates:
            status = "wrong_period"
        elif len(exact_candidates) > 1:
            status = "ambiguous"
        elif exact_candidates[0]["raw"] is None:
            status = "unresolved"
        else:
            status = "match" if _value(exact_candidates[0]["raw"]) == _value(raw) else "mismatch"
        matched_page = page_candidates[0]["page"] if page_candidates else page_number
        if page_candidates:
            matched_ir_pages[key] = matched_page
        details.append({"csv_line": csv_line, "statement": statement, "code": code,
                        "period_end": period.isoformat(), "period_type": period_type,
                        "pdf_page": matched_page, "printed_page": printed_pages.get(matched_page),
                        "reference_raw": raw, "ir_candidates": candidates, "status": status})

    statuses = ("match", "mismatch", "missing", "wrong_page", "wrong_period",
                "period_unresolved", "ambiguous", "unresolved", "invalid_reference")
    counts = {status: sum(item["status"] == status for item in details) for status in statuses}
    sha_present = "document_sha256" in gold_rows[0]
    if sha_present and any((row.get("document_sha256") or "").strip().lower()
                           != index["document_sha256"] for row in gold_rows):
        raise ValueError("CSV gold không cùng SHA-256 với PDF")
    review_methods = []
    review_issues = []
    for excel_row, row in enumerate(gold_rows, start=2):
        status = (row.get("review_status") or "").strip().lower()
        transcriber = (row.get("transcriber") or "").strip()
        checker = (row.get("checker") or "").strip()
        if status == "approved" and transcriber and checker and transcriber != checker:
            review_methods.append("independent_checked")
        elif status == "self_checked" and transcriber:
            review_methods.append("self_checked")
        else:
            review_methods.append("unreviewed")
            review_issues.append({"excel_row": excel_row, "transcriber": transcriber,
                                  "review_status": status})
    review_method = ("independent_checked" if review_methods and all(
        method == "independent_checked" for method in review_methods) else
        "self_checked" if review_methods and all(
            method in ("independent_checked", "self_checked") for method in review_methods) else
        "unreviewed")
    reviewed = review_method != "unreviewed"
    unreferenced: list[dict[str, Any]] = []
    period_unresolved_ir_cells = 0
    matched_candidate_count: dict[tuple[str, str, str, str], int] = {}
    for (statement, code, _), candidates in observed.items():
        for candidate in candidates:
            if candidate["period_end"] is None:
                period_unresolved_ir_cells += 1
                continue
            key = statement, code, candidate["period_end"], candidate["period_type"]
            if key in reference_pages and candidate["page"] == matched_ir_pages.get(key):
                matched_candidate_count[key] = matched_candidate_count.get(key, 0) + 1
                if matched_candidate_count[key] == 1:
                    continue
            unreferenced.append({"statement": statement, "code": code,
                                 "period_end": candidate["period_end"],
                                 "period_type": candidate["period_type"],
                                 "pdf_page": candidate["page"], "ir_raw": candidate["raw"]})
    return {"reference_file": str(gold_path.resolve()), "sheet": sheet_name,
            "reference_binding": "explicit_sha256" if sha_present else "pdf_via_ir_index",
            "items_total": len(details),
            "matched": counts["match"], "agreement_rate": counts["match"] / len(details),
            "human_approved": reviewed,
            "review_method": review_method,
            "review_issues": review_issues,
            "gold_accuracy": counts["match"] / len(details) if reviewed else None,
            "cells_correct": counts["match"],
            "cells_wrong_value": counts["mismatch"],
            "cells_missing": counts["missing"] + counts["wrong_page"] + counts["wrong_period"],
            "cells_extra": len(unreferenced),
            "status_counts": counts,
            "unreferenced_ir_cells": len(unreferenced),
            "unreferenced_ir_items": unreferenced,
            "period_unresolved_ir_cells": period_unresolved_ir_cells,
            "issues": [item for item in details if item["status"] != "match"],
            "per_item": details}


def _manual_gold_result(table_dir: Path, index: dict[str, Any], gold_path: Path) -> dict[str, Any]:
    """Compare independently transcribed values by statement, code and period."""
    gold = json.loads(gold_path.read_text(encoding="utf-8"))
    if gold.get("document_sha256") != index["document_sha256"]:
        raise ValueError("Manual gold không cùng SHA-256 với PDF")
    if gold.get("source") != "manual_pdf_transcription":
        raise ValueError("Manual gold phải được chép độc lập từ ảnh PDF")

    observed = _observed_cells(table_dir, index)

    seen: set[tuple[str, str, str]] = set()
    details: list[dict[str, Any]] = []
    for item in gold.get("items", []):
        key = item["statement"], str(item["code"]), str(item["period"])
        if key in seen:
            raise ValueError(f"Manual gold trùng khóa {key}")
        seen.add(key)
        if not item.get("reviewer") or not isinstance(item.get("gold_raw"), str):
            raise ValueError(f"Manual gold thiếu reviewer hoặc gold_raw tại {key}")
        candidates = observed.get(key, [])
        if not candidates:
            status = "missing"
        elif len({_value(candidate["raw"] or "") for candidate in candidates}) > 1:
            status = "ambiguous"
        elif candidates[0]["raw"] is None:
            status = "unresolved"
        else:
            status = "match" if _value(candidates[0]["raw"]) == _value(item["gold_raw"]) else "mismatch"
        details.append({"statement": key[0], "code": key[1], "period": key[2],
                        "gold_raw": item["gold_raw"], "ir_candidates": candidates, "status": status})
    matched = sum(item["status"] == "match" for item in details)
    review_methods = []
    for item in gold["items"]:
        status = str(item.get("review_status") or "").strip().lower()
        reviewer = str(item.get("reviewer") or "").strip()
        checker = str(item.get("checker") or "").strip()
        if status == "approved" and reviewer and checker and reviewer != checker:
            review_methods.append("independent_checked")
        elif status == "self_checked" and reviewer:
            review_methods.append("self_checked")
        else:
            review_methods.append("unreviewed")
    review_method = ("independent_checked" if review_methods and all(
        method == "independent_checked" for method in review_methods) else
        "self_checked" if review_methods and all(
            method in ("independent_checked", "self_checked") for method in review_methods) else
        "unreviewed")
    reviewed = review_method != "unreviewed"
    return {"source": "manual_pdf_transcription", "gold_file": str(gold_path),
            "gold_items": len(details), "gold_matched": matched,
            "human_approved": reviewed,
            "review_method": review_method,
            "agreement_rate": matched / len(details) if details else None,
            "gold_accuracy": matched / len(details) if reviewed else None,
            "status_counts": {status: sum(item["status"] == status for item in details)
                              for status in ("match", "mismatch", "missing", "ambiguous", "unresolved")},
            "per_item": details}


def evaluate(pdf: Path, ir_dir: Path, manual_gold: Path | None = None,
             gold_csv: Path | None = None, gold_xlsx: Path | None = None,
             gold_sheet: str | None = None, gold_only: bool = False) -> dict[str, Any]:
    table_dir = ir_dir / "table"
    index = json.loads((table_dir / "index.json").read_text(encoding="utf-8"))
    digest = hash_file(pdf)
    if index["document_sha256"] != digest:
        raise ValueError("PDF không trùng SHA-256 với table/index.json")
    if gold_only:
        if gold_csv is None and gold_xlsx is None and manual_gold is None:
            raise ValueError("--gold-only cần --gold-csv, --gold-xlsx hoặc --manual-gold")
        result: dict[str, Any] = {"method": "keyed_gold_comparison",
                                  "document_sha256": digest}
        if manual_gold is not None:
            result["manual_gold"] = _manual_gold_result(table_dir, index, manual_gold)
        if gold_csv is not None:
            result["csv_gold"] = _spreadsheet_gold_result(table_dir, index, gold_csv)
        if gold_xlsx is not None:
            result["xlsx_gold"] = _spreadsheet_gold_result(table_dir, index, gold_xlsx, gold_sheet)
        save_json(ir_dir / "cell_evaluation.json", result)
        return result
    rows: list[dict[str, Any]] = []
    for number in expand_page_ranges(index["structured_table_ranges"]):
        page = json.loads((table_dir / f"p-{number:03d}.ir.json").read_text(encoding="utf-8"))
        if page["source"]["engine"] != "TEXT_LAYER" or page["page_size"]["unit"] != "pt":
            continue
        words = _poppler_words(pdf, number)
        for block in page["blocks"]:
            if block["type"] != "table":
                continue
            for row in block["rows"]:
                for column, cell in row["cells"].items():
                    bbox = cell.get("bbox")
                    if bbox is None or cell.get("raw") is None:
                        continue
                    x0, y0, x1, y1 = bbox
                    selected = [word for word in words
                                if x0 - 2 <= (word["bbox"][0] + word["bbox"][2]) / 2 <= x1 + 2
                                and y0 - 2 <= (word["bbox"][1] + word["bbox"][3]) / 2 <= y1 + 2]
                    selected.sort(key=lambda word: (round(word["bbox"][1] / 3), word["bbox"][0]))
                    reference = " ".join(word["text"] for word in selected)
                    status = ("no_pdf_word" if not selected else
                              "match" if _value(reference) == _value(cell["raw"]) else "mismatch")
                    rows.append({"page": number, "table_id": block["id"], "row_id": row["id"],
                                 "code": row.get("code"), "column": column, "bbox": bbox,
                                 "ir_raw": cell["raw"], "pdf_raw": reference, "status": status})
    matched = sum(row["status"] == "match" for row in rows)
    result = {"method": "poppler_bbox_layout_tokenization_comparison",
              "document_sha256": digest,
              "cells_total": len(rows), "cells_matched": matched,
              "cells_mismatched": sum(row["status"] == "mismatch" for row in rows),
              "cells_without_pdf_word": sum(row["status"] == "no_pdf_word" for row in rows),
              "tokenization_agreement": matched / len(rows) if rows else None,
              "per_cell": rows}
    if manual_gold is not None:
        result["manual_gold"] = _manual_gold_result(table_dir, index, manual_gold)
    if gold_csv is not None:
        result["csv_gold"] = _spreadsheet_gold_result(table_dir, index, gold_csv)
    if gold_xlsx is not None:
        result["xlsx_gold"] = _spreadsheet_gold_result(table_dir, index, gold_xlsx, gold_sheet)
    save_json(ir_dir / "cell_evaluation.json", result)
    return result


def _print_summary(result: dict[str, Any], pdf: Path, ir_dir: Path) -> None:
    """Show the decision and actionable discrepancies; keep full evidence in JSON."""
    print(f"PDF: {pdf.name}")
    reference = result.get("xlsx_gold") or result.get("csv_gold")
    if reference is None:
        manual = result.get("manual_gold")
        if manual is not None and manual["agreement_rate"] is not None:
            label = ("Độ chính xác trên gold tự kiểm" if manual["review_method"] == "self_checked"
                     else "Độ chính xác trên gold kiểm độc lập" if manual["human_approved"]
                     else "Mức khớp với file chép tay (chưa duyệt)")
            print(f"{label}: {manual['agreement_rate'] * 100:.2f}% "
                  f"({manual['gold_matched']}/{manual['gold_items']} ô)")
        elif result.get("tokenization_agreement") is not None:
            print(f"Đồng thuận tách từ với Poppler: {result['tokenization_agreement'] * 100:.2f}% "
                  f"({result['cells_matched']}/{result['cells_total']} ô)")
        else:
            print("Không có ô có tọa độ để so với Poppler")
    else:
        label = ("Độ chính xác trên gold tự kiểm" if reference["review_method"] == "self_checked"
                 else "Độ chính xác trên gold kiểm độc lập" if reference["human_approved"]
                 else "Mức khớp với Excel/CSV (chưa duyệt)")
        print(f"{label}: {reference['agreement_rate'] * 100:.2f}% "
              f"({reference['matched']}/{reference['items_total']} ô)")
        issues = reference["issues"]
        print(f"Đúng: {reference['cells_correct']}; sai giá trị: {reference['cells_wrong_value']}; "
              f"thiếu/sai trang hoặc kỳ: {reference['cells_missing']}; "
              f"ô IR thừa: {reference['cells_extra']}")
        if reference["review_issues"]:
            print("Dòng Excel chưa xác nhận: " + ", ".join(
                str(item["excel_row"]) for item in reference["review_issues"][:20]))
            if len(reference["review_issues"]) > 20:
                print(f"... còn {len(reference['review_issues']) - 20} dòng")
        if len(issues) > reference["cells_wrong_value"] + reference["cells_missing"]:
            print(f"Cần giải quyết thêm: {len(issues) - reference['cells_wrong_value'] - reference['cells_missing']} ô")
        for issue in issues[:10]:
            printed = issue.get("printed_page")
            page = f"PDF {issue['pdf_page']}" + (f", số in {printed}" if printed else "")
            candidate = next((item["raw"] for item in issue["ir_candidates"]
                              if item["page"] == issue["pdf_page"]), None)
            print(f"  - {page}; {issue['statement']} mã {issue['code']}, "
                  f"kỳ {issue['period_end']}: {issue['status']}; "
                  f"tham chiếu={issue['reference_raw']!r}, IR={candidate!r}")
        if len(issues) > 10:
            print(f"  ... còn {len(issues) - 10} ô; xem file JSON để biết chi tiết")
    print(f"Chi tiết: {ir_dir / 'cell_evaluation.json'}")


def _evaluation_ok(result: dict[str, Any]) -> bool:
    reference_ok = all(key not in result or result[key]["matched"] == result[key]["items_total"]
                       for key in ("csv_gold", "xlsx_gold"))
    if "manual_gold" in result:
        reference_ok = (reference_ok and
                        result["manual_gold"]["gold_matched"] == result["manual_gold"]["gold_items"])
    poppler_ok = result.get("cells_mismatched", 0) == result.get("cells_without_pdf_word", 0) == 0
    return reference_ok and poppler_ok


def _cross_period_result(prior_ir_dir: Path, current_ir_dir: Path,
                         prior_year: int) -> dict[str, Any]:
    """Compare year N as originally reported with year N in the N+1 report.

    Differences need human review: a company may restate its comparative
    figures. This check does not identify which report, if either, is wrong.
    """
    def cells(ir_dir: Path) -> dict[tuple[str, str, str, str], list[dict[str, Any]]]:
        table_dir = ir_dir / "table"
        index = json.loads((table_dir / "index.json").read_text(encoding="utf-8"))
        result: dict[tuple[str, str, str, str], list[dict[str, Any]]] = {}
        for (statement, code, _), candidates in _observed_cells(table_dir, index).items():
            for candidate in candidates:
                period_end = candidate["period_end"]
                if period_end is None or not period_end.startswith(f"{prior_year}-"):
                    continue
                key = statement, code, period_end, candidate["period_type"]
                result.setdefault(key, []).append(candidate)
        return result

    original = cells(prior_ir_dir)
    comparative = cells(current_ir_dir)
    items: list[dict[str, Any]] = []
    for key in sorted(original.keys() | comparative.keys()):
        earlier, later = original.get(key, []), comparative.get(key, [])
        if not earlier:
            status = "missing_in_prior_report"
        elif not later:
            status = "missing_in_current_comparative"
        elif len(earlier) != 1 or len(later) != 1 or earlier[0]["raw"] is None or later[0]["raw"] is None:
            status = "ambiguous"
        else:
            status = ("match" if _value(earlier[0]["raw"]) == _value(later[0]["raw"])
                      else "CROSS_PERIOD_MISMATCH")
        items.append({"statement": key[0], "code": key[1], "period_end": key[2],
                      "period_type": key[3], "prior_report": earlier,
                      "current_report_comparative": later, "status": status})
    counts = {status: sum(item["status"] == status for item in items)
              for status in ("match", "CROSS_PERIOD_MISMATCH", "missing_in_prior_report",
                             "missing_in_current_comparative", "ambiguous")}
    return {"method": "same_period_cross_report_comparison",
            "period_year": prior_year, "counts": counts,
            "review_note": "Differences may be extraction errors or restated comparative figures.",
            "issues": [item for item in items if item["status"] != "match"]}


def _run_default_gold() -> int:
    """Run the project's configured gold sets with one command."""
    root = Path(__file__).resolve().parents[1]
    manifest = root / "Gold_Set" / "evaluations.json"
    payload = json.loads(manifest.read_text(encoding="utf-8"))
    jobs = payload.get("evaluations")
    if not isinstance(jobs, list) or not jobs:
        raise ValueError(f"Danh sách evaluations rỗng hoặc không hợp lệ: {manifest}")
    status = 0
    completed: dict[int, tuple[dict[str, Any], dict[str, Any]]] = {}
    for job in jobs:
        pdf = root / job["pdf"]
        ir_dir = root / job["ir_dir"]
        print(f"\n=== {pdf.name} ===")
        try:
            result = evaluate(pdf, ir_dir, gold_xlsx=root / job["gold_xlsx"],
                              gold_sheet=job["gold_sheet"], gold_only=True)
            report_year = int(job["report_year"])
            if report_year - 1 in completed:
                prior_job, _ = completed[report_year - 1]
                result["cross_period"] = _cross_period_result(
                    root / prior_job["ir_dir"], ir_dir, report_year - 1)
                counts = result["cross_period"]["counts"]
                print(f"So kỳ {report_year - 1} giữa hai bản: "
                      f"{counts['match']} khớp, "
                      f"{counts['CROSS_PERIOD_MISMATCH']} lệch; "
                      f"xem chi tiết trong kết quả JSON")
            save_json(root / job["result_json"], result)
            completed[report_year] = job, result
            _print_summary(result, pdf, ir_dir)
            if not _evaluation_ok(result):
                status = max(status, 1)
        except (OSError, ValueError, KeyError) as exc:
            print(f"Không đánh giá được {pdf.name}: {exc}", file=sys.stderr)
            status = 2
    return status


def main() -> int:
    if sys.stdout.encoding.lower() != "utf-8":
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pdf", type=Path, nargs="?",
                        help="Bỏ trống để tự đánh giá các gold set trong Gold_Set/evaluations.json")
    parser.add_argument("--ir-dir", type=Path)
    parser.add_argument("--manual-gold", type=Path,
                        help="JSON chép tay độc lập, khóa theo statement/code/period")
    gold_options = parser.add_mutually_exclusive_group()
    gold_options.add_argument("--gold-csv", type=Path,
                        help="CSV statement/code/period_end/period_type/value_raw/pdf_page; đo theo dòng/cột")
    gold_options.add_argument("--gold-xlsx", type=Path,
                              help="Excel gold set; dùng cùng --gold-sheet để chọn năm")
    parser.add_argument("--gold-sheet", help="Tên sheet trong Excel, ví dụ BCTN_2024")
    parser.add_argument("--gold-only", action="store_true",
                        help="Chỉ so theo gold set, không cần pdftotext/Poppler")
    parser.add_argument("--result-json", type=Path,
                        help="Lưu thêm bản sao kết quả tại đường dẫn này")
    args = parser.parse_args()
    if args.pdf is None:
        if any((args.ir_dir, args.manual_gold, args.gold_csv, args.gold_xlsx,
                args.gold_sheet, args.gold_only, args.result_json)):
            parser.error("Các tùy chọn đánh giá riêng cần truyền đường dẫn PDF")
        return _run_default_gold()
    pdf = args.pdf.resolve()
    ir_dir = args.ir_dir or Path("ocr_router_out") / pdf.stem / "annual_ir"
    result = evaluate(pdf, ir_dir, args.manual_gold, args.gold_csv,
                      args.gold_xlsx, args.gold_sheet, args.gold_only)
    if args.result_json is not None:
        save_json(args.result_json, result)
    _print_summary(result, pdf, ir_dir)
    return 0 if _evaluation_ok(result) else 1


if __name__ == "__main__":
    raise SystemExit(main())
